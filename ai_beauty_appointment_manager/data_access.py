from pathlib import Path
from openpyxl import load_workbook


DATA_FILE = Path(__file__).resolve().parent.parent / "data" / "business_data.xlsx"


def load_schedule():
   
    workbook = load_workbook(
        DATA_FILE,
        read_only=True,
        data_only=True
    )

    sheet = workbook["DemoSchedule"]

    schedule = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        schedule[row[0]] = row[1]

    workbook.close()
    
    return schedule

def load_service_prices():
    workbook = load_workbook(
        DATA_FILE,
        read_only=True,
        data_only=True
    )

    sheet = workbook["Services"]

    prices = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[4] is True:
            prices[row[1]] = row[2]

    workbook.close()
    
    return prices


if __name__ == "__main__":
    print(load_schedule())
    print(load_service_prices())
    